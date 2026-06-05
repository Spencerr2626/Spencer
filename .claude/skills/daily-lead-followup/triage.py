#!/usr/bin/env python3
"""Daily lead follow-up triage for Spencer's Kommons CRM.

Reads the CRM workbook, finds follow-ups that are due/overdue, sorts them by
priority (Hot > Warm > Cold) then by date, and prints a structured report that
the assistant uses to draft per-lead messages.

Usage:
    python3 triage.py [--date YYYY-MM-DD] [--horizon DAYS] [--crm PATH]

--date     "today" to triage against (default: today in Asia/Kuala_Lumpur, GMT+8)
--horizon  also include follow-ups due within this many days ahead (default: 2)
--crm      path to the workbook (default: references/crm/Spencer_Kommons_CRM_v3.xlsx)
"""
import argparse
import datetime as dt
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl not installed. Run: pip install openpyxl")

PRIORITY_RANK = {"hot": 0, "warm": 1, "cold": 2}


def malaysia_today():
    return (dt.datetime.utcnow() + dt.timedelta(hours=8)).date()


def parse_date(v):
    """Parse a cell that may be a datetime, a DD/MM/YYYY string, or junk."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None  # unparseable (e.g. "7/6/2026 TBC")


def priority_rank(tag):
    t = (tag or "").lower()
    for key, rank in PRIORITY_RANK.items():
        if key in t:
            return rank
    return 3


def clean_priority(tag):
    t = (tag or "").lower()
    for key in PRIORITY_RANK:
        if key in t:
            return key.capitalize()
    return "Unknown"


def sheet_records(ws, header_row=3):
    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    headers = [h for h in headers if h]
    n = len(headers)
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, n + 1)]
        if not any(v not in (None, "") for v in vals):
            continue
        out.append(dict(zip(headers, vals)))
    return out


def is_done(rec):
    d = str(rec.get("Done?", "") or "").strip().lower()
    return d in ("done", "yes", "y", "✓", "true")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date")
    ap.add_argument("--horizon", type=int, default=2)
    ap.add_argument("--crm", default="references/crm/Spencer_Kommons_CRM_v3.xlsx")
    args = ap.parse_args()

    today = parse_date(args.date) if args.date else malaysia_today()
    if today is None:
        sys.exit(f"Could not parse --date '{args.date}'")
    horizon = today + dt.timedelta(days=args.horizon)

    path = Path(args.crm)
    if not path.exists():
        sys.exit(f"CRM not found at {path}")
    wb = openpyxl.load_workbook(path, data_only=True)

    def find_sheet(keyword):
        for name in wb.sheetnames:
            if keyword.lower() in name.lower():
                return wb[name]
        return None

    fu = find_sheet("Follow-Up Schedule")
    al = find_sheet("Active Leads")

    # Index active leads by ID for enrichment (language, contact, area, budget...)
    leads = {}
    if al:
        for rec in sheet_records(al):
            lid = rec.get("ID")
            if lid:
                leads[str(lid).strip()] = rec

    items = []
    if fu:
        for rec in sheet_records(fu):
            if is_done(rec):
                continue
            d = parse_date(rec.get("Follow-Up Date"))
            if d is None or d > horizon:
                continue
            lid = str(rec.get("ID") or "").strip()
            lead = leads.get(lid, {})
            items.append({
                "id": lid,
                "name": rec.get("Contact Name") or lead.get("Name"),
                "date": d,
                "priority": clean_priority(rec.get("Priority")),
                "prank": priority_rank(rec.get("Priority")),
                "channel": rec.get("Channel"),
                "language": rec.get("Language") or lead.get("Language"),
                "nationality": rec.get("Nationality") or lead.get("Nationality"),
                "purpose": rec.get("Purpose / Topic"),
                "next_action": rec.get("Next Action"),
                "outcome": rec.get("Outcome / Notes"),
                "appt": rec.get("Appointment Date & Time"),
                "remarks": rec.get("Remarks"),
                # extra context from Active Leads
                "area": lead.get("Area"),
                "budget": lead.get("Budget (RM)"),
                "lead_type": lead.get("Lead Type"),
                "status": lead.get("Status"),
                "wechat": lead.get("WeChat ID"),
                "notes": lead.get("Notes"),
            })

    items.sort(key=lambda x: (x["prank"], x["date"] or dt.date.max))

    overdue = [i for i in items if i["date"] and i["date"] < today]
    due_today = [i for i in items if i["date"] == today]
    upcoming = [i for i in items if i["date"] and i["date"] > today]

    def bucket(label, rows):
        print(f"\n### {label} ({len(rows)})")
        if not rows:
            print("  (none)")
        for i in rows:
            budget = f" | budget RM{int(i['budget']):,}" if isinstance(i["budget"], (int, float)) else ""
            print(f"- [{i['id']}] {i['name']} — {i['priority']} | {i['date'].strftime('%d/%m/%Y')} | "
                  f"{i['channel']} | {i['language']}{budget}")
            if i["purpose"]:
                print(f"    Purpose: {i['purpose']}")
            if i["next_action"]:
                print(f"    Next action: {i['next_action']}")
            if i["appt"]:
                print(f"    Appt: {i['appt']}")
            ctx = i.get("notes") or i.get("outcome")
            if ctx:
                print(f"    Context: {ctx}")

    print(f"# Daily Follow-Up Triage — {today.strftime('%A %d %b %Y')} (GMT+8)")
    print(f"Workbook: {path.name} | Horizon: +{args.horizon}d")
    print(f"Totals: {len(overdue)} overdue, {len(due_today)} due today, {len(upcoming)} upcoming (≤+{args.horizon}d)")

    bucket("⚠️ OVERDUE", overdue)
    bucket("📌 DUE TODAY", due_today)
    bucket("🔜 UPCOMING", upcoming)


if __name__ == "__main__":
    main()
