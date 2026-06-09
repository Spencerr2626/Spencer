#!/usr/bin/env python3
"""
Lease Radar — surfaces tenancy RENEWALS (lease ending soon) and CHECK-INS
(lease starting / move-in soon) from the Kommons CRM, so Spencer never
misses a renewal window or a move-in.

Reads:
  • "✅ Done Deals" — Lease Start / Lease End per tenant
  • "💰 Pipeline & Deals" — VP / Move-in for sale completions

Usage:
    python lease_radar.py [--date YYYY-MM-DD] [--renewal-days 60] [--checkin-days 30]

Defaults: today in GMT+8, renewals within 60 days, check-ins within 30 days.
Also flags leases that ended in the last 45 days (confirm renewed/vacated).
"""
import argparse
import sys
from datetime import datetime, date, timedelta, timezone

try:
    import openpyxl
except ImportError:
    sys.exit("Needs openpyxl:  pip install openpyxl")

CRM = "references/crm/Spencer_Kommons_CRM_v3.xlsx"


def parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for f in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v.strip(), f).date()
            except ValueError:
                pass
    return None


def malaysia_today():
    return datetime.now(timezone(timedelta(hours=8))).date()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date")
    ap.add_argument("--renewal-days", type=int, default=60)
    ap.add_argument("--checkin-days", type=int, default=30)
    ap.add_argument("--crm", default=CRM)
    args = ap.parse_args()

    today = parse_date(args.date) if args.date else malaysia_today()
    wb = openpyxl.load_workbook(args.crm, data_only=True)

    renewals, checkins, expired = [], [], []

    dd = wb["✅ Done Deals"]
    for r in range(4, dd.max_row + 1):
        name = dd.cell(r, 2).value
        if not name:
            continue
        prop = dd.cell(r, 16).value or ""
        unit = dd.cell(r, 17).value or ""
        where = f"{prop} {unit}".strip()
        ls = parse_date(dd.cell(r, 18).value)
        le = parse_date(dd.cell(r, 19).value)
        if le:
            d = (le - today).days
            if 0 <= d <= args.renewal_days:
                renewals.append((d, le, name, where, "renewal"))
            elif -45 <= d < 0:
                expired.append((d, le, name, where, "expired"))
        if ls:
            d = (ls - today).days
            if -3 <= d <= args.checkin_days:
                checkins.append((d, ls, name, where, "check-in"))

    pl = wb["💰 Pipeline & Deals"]
    for r in range(4, pl.max_row + 1):
        name = pl.cell(r, 2).value
        if not name:
            continue
        proj = pl.cell(r, 4).value or ""
        vp = parse_date(pl.cell(r, 13).value)
        if vp:
            d = (vp - today).days
            if -3 <= d <= args.checkin_days:
                checkins.append((d, vp, name, proj, "move-in (sale)"))

    renewals.sort()
    checkins.sort()
    expired.sort()

    print(f"# Lease Radar — {today:%a %d %b %Y} (GMT+8)")
    print(f"Renewals ≤{args.renewal_days}d · Check-ins ≤{args.checkin_days}d\n")

    def line(d, dt, name, where, kind):
        when = f"in {d}d" if d > 0 else ("TODAY" if d == 0 else f"{-d}d ago")
        return f"- **{name}** — {where} · {dt:%d/%m/%Y} ({when})"

    if expired:
        print("### ⚠️ JUST EXPIRED — confirm renewed or vacated")
        for x in expired:
            print(line(*x))
        print()
    if renewals:
        print(f"### 🔁 RENEWALS DUE ({len(renewals)})")
        for x in renewals:
            print(line(*x))
        print()
    if checkins:
        print(f"### 🔑 CHECK-INS / MOVE-INS ({len(checkins)})")
        for x in checkins:
            print(line(*x))
        print()
    if not (expired or renewals or checkins):
        print("Nothing on the radar in the window. ✅")


if __name__ == "__main__":
    main()
